"""Streaming XLSX / CSV inspection. No source financial semantics are inferred."""
import csv, re, zipfile, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from openpyxl import load_workbook
from fastapi import HTTPException

ALIASES = {
 'uhid':['uhid','patient id'], 'admission_id':['ipd','file ipd','file id','file no','admission id'],
 'patient':['patient name','patient'], 'panel':['panel','tpa'], 'branch':['branch','hospital branch'],
 'admission_date':['admission date','admit date'], 'discharge_date':['discharge date'],
 'date':['transaction date','date','sale date','service date'], 'bill_id':['bill no','bill number'],
 'description':['service description','service name','medicine name','item name','description'],
 'code':['service code','item code','medicine code'], 'quantity':['quantity','qty'],
 'status':['status'], 'classification':['patient type','opd/ipd'],
}

def norm(s): return re.sub(r'\s+', ' ', str(s or '').strip()).casefold()
def identifier(v):
    if v is None: return None
    return str(v).strip() or None

def cell_value(cell):
    v=cell.value
    if isinstance(v,(datetime.date,datetime.datetime)): return v.isoformat()[:10]
    if v is None: return ''
    if isinstance(v,(int,float)):
        fmt=cell.number_format or ''
        if re.fullmatch(r'0+',fmt) and float(v).is_integer(): return str(int(v)).zfill(len(fmt))
        if float(v).is_integer(): return str(int(v))
    return str(v)

def sheets(path):
    if str(path).lower().endswith('.csv'):
        with open(path,encoding='utf-8-sig',newline='') as f:
            yield 'CSV', enumerate(csv.reader(f),1)
    else:
        # Reject archive bombs before openpyxl expands workbook XML.
        with zipfile.ZipFile(path) as z:
            if sum(x.file_size for x in z.infolist())>250_000_000:
                raise HTTPException(400,'Workbook expanded size exceeds 250 MB.')
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            for ws in wb.worksheets:
                ws.reset_dimensions()
                yield ws.title, ((i,[cell_value(c) for c in row]) for i,row in enumerate(ws.iter_rows(),1))
        finally: wb.close()

def inspect(path):
    result=[]
    for name, rows in sheets(path):
        nonblank=[]; count=0; candidates=[]
        for i,row in rows:
            if not any(str(x).strip() for x in row): continue
            count+=1
            if len(nonblank)<18: nonblank.append({'row':i,'cells':row[:80]})
            if count<=100:
                score=sum(any(norm(v) in vals for vals in ALIASES.values()) for v in row)
                if score: candidates.append((score,i,row))
        best=max(candidates,default=(0,1,[]),key=lambda v:v[0])
        suggestions={f:str(next(j for j,v in enumerate(best[2]) if norm(v) in aliases)) for f,aliases in ALIASES.items() if any(norm(v) in aliases for v in best[2])}
        result.append({'name':name,'nonblank_rows':count,'sample':nonblank,'header_row_candidate':best[1], 'columns':best[2], 'suggestions':suggestions})
    return result

def decimal_value(v, precision='.01'):
    s=str('' if v is None else v).strip().replace('₹','').replace(',','')
    if not s or s in ('-','N/A','NA'): return None
    if s.startswith('(') and s.endswith(')'): s='-'+s[1:-1]
    try:
        n=Decimal(s)
        if not n.is_finite(): raise ValueError('Non-finite amount')
        return format(n.quantize(Decimal(precision),rounding=ROUND_HALF_UP),'f')
    except (InvalidOperation,ValueError): raise ValueError('Invalid numeric amount')

def date_value(v):
    if not str(v or '').strip(): return None
    s=str(v).strip()
    for fmt in ('%Y-%m-%d','%Y-%m-%dT%H:%M:%S','%d/%m/%Y','%d-%m-%Y','%d.%m.%Y','%d/%m/%y','%d-%b-%Y','%d %b %Y'):
        try: return datetime.datetime.strptime(s,fmt).date().isoformat()
        except ValueError: pass
    raise ValueError('Unrecognised date; use ISO or Indian day/month/year.')
